# -*- coding: utf-8 -*-
"""主流程：读源 .xls -> 抽取内嵌 PDF -> 解析 -> 填模板 -> 输出。

产出（均在用户指定的输出目录，全平铺）：
    {prefix}-{材料名}.xlsx      每个 K 列合并格一份（重名自动加 -1 / -2）
    HSF_{prefix}-{材料名}.pdf   该格所有 HSF 报告合并成一份
    MSDS_{prefix}-{材料名}.pdf  每个 K 列合并格一份，与 HSF 一一对应
    processing_log.txt          处理日志
    _汇总报告.xlsx              汇总表（告警标红）

2026-08-31 v1.3：MSDS 由「按 M 列合并格另存」改为「跟随 K 块」。
一个 M 列合并格常常横跨多个 K 块（如 RR12346 行 32-37 覆盖 5 个 K 块，
却只有 1 份 MSDS），旧逻辑只出 1 份，导致 MSDS 与 HSF 对不上号。
现在每个 K 块都按自己所属 M 格取同一份 PDF 另存，内容相同、名称各异，
份数与 HSF 严格一致。

2026-10-12 v1.4：
  1. 报告图标归属改按**底边**所在行（Settings.pdf_row_anchor），
     压在两格之间的图标不再错挂，且接收方/失去方都写告警。
  2. 同一个 M 格内 D 列名称相同、G 列 CAS 也相同的多个 K 块**合并成一块**：
     重量相加后重新归一化到 100%，HSF 报告全并进来、几个编号写几行。
     合并发生在空块剔除之前，否则「自己那格没报告」的行会被先丢掉。
"""
from __future__ import annotations

import os
import traceback
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Font, PatternFill

from .config import (DOMAIN_GROUPS, FALLBACK_REPORT_NO, MANDATORY_ITEMS,
                     RESULT_ITEMS, SECRET_CAS_VALUE, Settings)
from .ole_extract import EmbeddedObject, extract_embedded_objects, is_pdf
from .ocr import ocr_available
from .pdf_parse import ReportInfo, merge_results, parse_pdf
from .pdfout import save_pdf_group
from .util import (clean_cas, extract_prefix, material_display, remove_spaces,
                   safe_name)
from .xls_source import (Component, MaterialBlock, SourceSheet, merge_components,
                         normalize_content)
from .xlsx_build import (ComponentRow, MaterialRecord, ReportRow, build_workbook,
                         _is_secret)

RED_FILL = PatternFill('solid', fgColor='FFC7CE')
BOLD = Font(bold=True)


# ---------------------------------------------------------------- 名称

def _material_name(sheet: SourceSheet, block: MaterialBlock, st: Settings) -> str:
    """均质名称主体（不含 RR 前缀）。

    2026-08-31 需求：D 列里的括号内容不取，如 Die1(EMMJR68) -> Die1。
    2026-10-12：实现挪到 SourceSheet.material_name——build_blocks 里也要用
    同一个名称做「同名同 CAS 合并」的键，两处必须是同一套口径。
    """
    return sheet.material_name(block)


def _dedup(names: list[str]) -> list[str]:
    """重名加 -1 / -2 后缀（首次出现不带后缀）。"""
    seen: dict = {}
    out = []
    for n in names:
        seen[n] = seen.get(n, 0) + 1
        out.append(n if seen[n] == 1 else f'{n}-{seen[n] - 1}')
    return out


# ---------------------------------------------------------------- 结果补全

def apply_rules(results: dict, rec_warns: list):
    """领域规则推断 + 必填项兜底。返回被补成 ND 的项目列表。"""
    inferred: list = []
    for group in DOMAIN_GROUPS:
        if any(results.get(i) for i in group):
            for i in group:
                if not results.get(i):
                    results[i] = 'ND'
                    inferred.append(i)
    for i in MANDATORY_ITEMS:
        if not results.get(i):
            results[i] = 'ND'
            inferred.append(i)
            rec_warns.append(f'{i} 未在任何报告中测出，已按规则补 ND')
    return inferred


# ---------------------------------------------------------------- 主流程

@dataclass
class RunResult:
    out_dir: str = ''
    records: list = field(default_factory=list)
    logs: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    skipped: list = field(default_factory=list)   # [(row_start, row_end, reason)]
    n_xlsx: int = 0
    n_hsf: int = 0
    n_msds: int = 0


class Pipeline:
    def __init__(self, src_path: str, tpl_path: str, out_dir: str,
                 settings: Settings | None = None, progress=None, logger=None,
                 prefix_resolver=None, prefix_override: str | None = None):
        self.src_path = src_path
        self.tpl_path = tpl_path
        self.out_dir = out_dir
        self.st = settings or Settings()
        self.progress = progress or (lambda msg, pct: None)
        self.logger = logger or (lambda line, level: None)
        # 抓不到 RR 编号时的兜底顺序：调用方已问到的值 → 回调（如弹窗）→ 中止
        # 2026-08-31：去掉硬编码 fallback，抓不到就一定要问用户
        self.prefix_resolver = prefix_resolver
        self.prefix_override = (prefix_override or '').strip() or None
        self.res = RunResult(out_dir=out_dir)
        self._log_lines: list = []
        # 本次运行已占用的 PDF 路径。只用于批内避重，不用它探测磁盘——
        # 探测磁盘会让重跑到同一目录时产物翻倍（详见 pdfout.save_pdf_group 注释）。
        self._used_pdf_paths: set = set()

    # ---------------- 日志
    def log(self, msg: str, level: str = 'INFO'):
        line = f'[{level}] {msg}'
        self._log_lines.append(line)
        self.res.logs.append(line)
        try:
            self.logger(line, level)
        except Exception:
            pass

    # ---------------- 芯片编号
    def _resolve_prefix(self) -> str:
        """确定芯片编号（RR12345 这类），抓不到就问用户。

        三级来源：① 源文件名自动抓取 ② 调用方已问到的值（CLI --prefix / GUI 弹窗）
        ③ 调用方回调。都没有则返回空串，由 run() 中止。
        2026-08-31 之前是硬编码兜底成 'RR12345'，换芯片时会静默套错编号，已去掉。
        """
        got = extract_prefix(self.src_path)
        if got:
            self.log(f'芯片编号从源文件名自动识别：{got}')
            return got

        fname = os.path.basename(self.src_path)
        self.log(f'源文件名中未找到 RR 编号：{fname}', 'WARN')

        if self.prefix_override:
            self.log(f'使用指定编号：{self.prefix_override}')
            return self.prefix_override

        if self.prefix_resolver:
            try:
                asked = (self.prefix_resolver(fname) or '').strip()
            except Exception as e:
                self.log(f'编号输入被取消或出错：{e}', 'WARN')
                asked = ''
            if asked:
                self.log(f'使用用户输入的编号：{asked}')
                return asked
            self.log('用户未输入编号', 'WARN')

        return ''

    # ---------------- 执行
    def run(self) -> RunResult:
        st = self.st
        try:
            os.makedirs(self.out_dir, exist_ok=True)
        except Exception as e:
            self.res.errors.append(f'输出目录不可创建：{e}')
            return self.res

        prefix = self._resolve_prefix()
        if not prefix:
            self.res.errors.append(
                '未能确定芯片编号（源文件名中找不到 RR 编号，且未提供有效输入），已中止')
            return self.res

        self.log(f'源表：{self.src_path}')
        self.log(f'模板：{self.tpl_path}')
        self.log(f'输出目录：{self.out_dir}')
        self.log(f'芯片编号前缀：{prefix}')

        # OCR 是可选依赖，精简版刻意不打进去。但缺了它，扫描件和「文本层解析
        # 失败」的报告就会静默退化成 TBD00000000 / NA —— 用户只看到数据变少，
        # 却不知道是 OCR 没装。静默降级比报错更危险，这里必须显式告警。
        if not ocr_available():
            self.log('当前程序未集成 OCR（精简版）：扫描件、以及常规解析失败的'
                     '报告，其编号会填 TBD00000000、测试结果会填 NA，需人工补录。'
                     '如需自动识别，请改用完整版', 'WARN')

        self.progress('读取源表并抽取内嵌 PDF', 5)
        try:
            objs = extract_embedded_objects(self.src_path)
        except Exception:
            self.log(traceback.format_exc(), 'ERROR')
            self.res.errors.append('源 .xls 解析失败（可能不是 BIFF8 格式或文件损坏）')
            return self.res
        pdfs = [o for o in objs if is_pdf(o)]
        self.log(f'共发现内嵌对象 {len(objs)} 个，其中 PDF {len(pdfs)} 个')
        if not pdfs:
            self.res.errors.append('源表中没有内嵌 PDF，无法生成结果')
            return self.res

        try:
            sheet = SourceSheet(self.src_path, st)
        except Exception:
            self.log(traceback.format_exc(), 'ERROR')
            self.res.errors.append('源 .xls 工作表读取失败')
            return self.res

        # M 列合并格清单：既用于「同名同 CAS 合并」的前提判断，也用于给每个
        # 材料块取 MSDS。这里必须包含没有内嵌 MSDS 的格子——只要行落在同一个
        # M 格里就该合并，缺 MSDS 是另一回事。
        r0, r1 = sheet.detect_data_range()
        m_cells_all = sheet.distinct_cells(st.columns.msds, r0, r1)

        def m_anchor(row: int):
            for (rs, re_) in m_cells_all:
                if rs <= row <= re_:
                    return rs
            return None

        try:
            blocks, skipped = sheet.build_blocks(objs, m_anchor=m_anchor)
        except Exception:
            # 物料块划分/合并里的任何异常都不该让进程炸掉——GUI 下只会看到
            # 一个空输出目录。宁可给出明确错误，也不要静默半成品。
            self.log(traceback.format_exc(), 'ERROR')
            self.res.errors.append('物料块划分失败（列位配置或源表结构异常）')
            return self.res
        self.res.skipped = skipped
        n_cells = sum(len(b.merged_rows) or 1 for b in blocks) + len(skipped)
        n_merged = sum(1 for b in blocks if b.is_merged)
        self.log(f'K 列合并格 {n_cells} 个：有效 {len(blocks)} 个'
                 + (f'（其中 {n_merged} 个由多个格合并而成）' if n_merged else '')
                 + f'，空块 {len(skipped)} 个')
        for rs, re_, reason in skipped:
            self.log(f'跳过空块 行 {rs + 1}-{re_ + 1}（{reason}），不出表', 'WARN')
        for b in blocks:
            if not b.is_merged:
                continue
            self.log(f'同名同 CAS 合并：源表行 {b.excel_rows} 同属一个 M 列合并格，'
                     f'D 列名称与 G 列 CAS 完全一致，已合并为一块——'
                     f'重量相加后重新归一化到 100%，{len(b.hsf_objs)} 份 HSF 报告'
                     f'合并成一份（报告编号照写 {len(b.hsf_objs)} 个）')
        self.log(f'将生成 {len(blocks)} 份表')
        if not blocks:
            self.res.errors.append(
                f'未识别到任何有效物料块（{n_cells} 个合并格全为空块），请检查列位配置')
            return self.res

        raw_names = [_material_name(sheet, b, st) for b in blocks]
        disp_names = _dedup(raw_names)

        # M 列合并格 -> 该格内嵌的 MSDS PDF。K 块按行号查表找到自己那份，
        # 从而实现「每个 K 块一份 MSDS」（2026-08-31 v1.3）。
        # 2026-10-12：M 格清单在 build_blocks 之前就要算好——它同时是
        # 「同名同 CAS 合并」的前提条件（必须同一个 M 格内才合并）。
        msds_cells = []
        for (rs, re_) in m_cells_all:
            group = [o for o in pdfs if o.col == st.columns.msds
                     and rs <= o.row <= re_]
            if group:
                msds_cells.append((rs, re_, group))
        self.log(f'M 列合并格 {len(msds_cells)} 个含 MSDS，'
                 f'将按 {len(blocks)} 个材料块各存一份')

        total = len(blocks) + 2
        for i, (block, disp) in enumerate(zip(blocks, disp_names)):
            pct = 10 + int(80 * i / max(total, 1))
            self.progress(f'处理 {i + 1}/{len(blocks)}：{disp}', pct)
            try:
                rec = self._one(sheet, block, disp, prefix, i, msds_cells)
            except Exception:
                self.log(traceback.format_exc(), 'ERROR')
                rec = MaterialRecord(index=i, name=disp, display=disp,
                                     excel_rows=block.excel_rows)
                rec.warnings.append('处理异常，输出表可能不完整')
            self.res.records.append(rec)

        self.progress('写日志与汇总报告', 96)
        self._write_log()
        self._write_summary(prefix)
        return self.res

    # ---------------- 单个物料
    def _one(self, sheet: SourceSheet, block: MaterialBlock, disp: str,
             prefix: str, idx: int, msds_cells: list) -> MaterialRecord:
        st = self.st
        # 2026-08-31：B2 均质名称要带芯片编号前缀（RR12345-Die1），与文件名一致。
        # 之前只填了 disp（Die1），漏了前缀——回读 B2 才发现。
        rec = MaterialRecord(index=idx, name=remove_spaces(f'{prefix}-{disp}'),
                             display=disp, excel_rows=block.excel_rows)
        warns = rec.warnings

        # 成分：先按 CAS 合并，再归一化
        # 2026-10-12：合并块（同名同 CAS 的多个 K 格）在这里天然处理正确——
        # 各成员成分行拼在一起后，merge_components 会把相同 CAS 的重量相加，
        # normalize_content 再把合并后的整体归一化到 100%。
        comps, merge_notes = merge_components(block.components)
        warns.extend(merge_notes)
        if block.is_merged:
            by_cas: dict = {}
            for c in block.components:
                k = clean_cas(c.cas) or f'{c.substance or "未命名"}(G{c.row + 1})'
                by_cas[k] = by_cas.get(k, 0.0) + max(c.mass, 0.0)
            detail = '、'.join(f'{k} {v:.6g}' for k, v in by_cas.items())
            self.log(f'{disp} 合并重量明细（H 列相加）：{detail}')
            warns.append(
                f'本表由源表 {block.excel_rows} 行合并而成（同一个 M 列合并格内'
                f'D 列名称与 G 列 CAS 完全相同）：重量相加后重新归一化为 100%，'
                f'{len(block.hsf_objs)} 份 HSF 报告合并成一份、报告编号写 '
                f'{len(block.hsf_objs)} 个')
        warns.extend(block.attr_notes)
        contents = normalize_content(comps, st.content_decimals)
        n_cleaned = 0
        for c, pct_v in zip(comps, contents):
            raw = (c.cas or '').strip()
            cas = clean_cas(raw)          # 只留数字和 '-'（2026-08-31 需求）
            if raw and cas and raw != cas:
                n_cleaned += 1
            # 保密判定看清洗后的值：'Trade secret' / 'Proprietary' 清洗后为空；
            # 只剩 '-' 这类没有数字的残留也按保密处理，避免把 '--' 写进 CAS 栏。
            secret = (_is_secret(raw, st.secret_when_cas_empty)
                      or not any(ch.isdigit() for ch in cas))
            rec.components.append(ComponentRow(
                substance=(c.substance or '').strip(),
                cas=cas,
                content=pct_v,
                secret=secret))
        if not rec.components:
            warns.append('未识别到成分行')
        n_sec = sum(1 for c in rec.components if c.secret)
        if n_sec:
            warns.append(f'存在 {n_sec} 条无 CAS NO 的成分，已按保密处理'
                         f'（是否保密=是，CAS NO={SECRET_CAS_VALUE}）')
        if n_cleaned:
            warns.append(f'{n_cleaned} 条 CAS NO 含数字和 - 以外的字符，已清洗'
                         f'（如 7440-57-5(金) -> 7440-57-5）')
        if any(not (c.substance or '').strip() for c in rec.components):
            warns.append('存在物质名称为空的成分行，请人工核对源表')

        # HSF 报告
        infos: list[ReportInfo] = []
        for o in block.hsf_objs:
            try:
                info = parse_pdf(o.data, o.file_name, o.file_name)
            except Exception as e:
                # 2026-08-31：单份 PDF 解析失败不再炸掉整块。
                # 之前一份加密/损坏的 PDF 会让整个材料的 xlsx 都不生成——
                # RR12346 的 AuPlating 就是这么丢的，而且日志里只有一串
                # traceback，看不出到底是哪份报告出的问题。
                warns.append(f'{o.file_name}：解析失败已跳过'
                             f'（{type(e).__name__}: {e}）')
                self.log(f'{disp}：{o.file_name} 解析失败，已跳过：'
                         f'{type(e).__name__}: {e}', 'ERROR')
                continue
            infos.append(info)
            if info.warnings:
                warns.append(f'{o.file_name}：' + '；'.join(info.warnings))

        merged, conflict = merge_results(infos)
        rec.conflicts = sorted(conflict)
        if conflict:
            warns.append('多份报告取值不一致（已取最不利值）：'
                         + '、'.join(sorted(conflict)))
        rec.inferred = apply_rules(merged, warns)
        rec.results = merged

        # 报告编号/日期兜底（2026-08-31 需求）
        known_dates = [i.report_date for i in infos if i.report_date]
        borrow_date = known_dates[0] if known_dates else ''
        n_no_fb, n_date_fb, n_date_missing = 0, 0, 0
        for info in infos:
            no = (info.report_no or '').strip()
            if not no:
                no = FALLBACK_REPORT_NO
                n_no_fb += 1
            date = (info.report_date or '').strip()
            if not date:
                if borrow_date:
                    # 同单元格多张报告日期本应一致，识别不出时借用同格其它报告的
                    date = borrow_date
                    n_date_fb += 1
                else:
                    n_date_missing += 1
            rec.reports.append(ReportRow(agency=info.agency, report_no=no,
                                         report_date=date))
        if n_no_fb:
            warns.append(f'{n_no_fb} 份报告未识别到编号，已填占位值 '
                         f'{FALLBACK_REPORT_NO}')
        if n_date_fb:
            warns.append(f'{n_date_fb} 份报告未识别到日期，已借用同单元格其它'
                         f'报告的日期 {borrow_date}')
        if n_date_missing:
            # grill-me 决策：整个单元格都识别不出时留空 + 标红，不编造日期
            warns.append(f'{n_date_missing} 份报告未识别到日期，且同单元格无可用'
                         f'日期可借用，日期栏留空待人工补')
        if not block.hsf_objs:
            warns.append('#该 K 列单元格内没有内嵌 HSF 报告，报告区留空待补')

        # 输出 xlsx
        fname = safe_name(f'{prefix}-{disp}', default=f'{prefix}-{idx}') + '.xlsx'
        out_path = os.path.join(self.out_dir, fname)
        warns.extend(build_workbook(self.tpl_path, out_path, rec, st))
        rec.out_xlsx = out_path
        self.res.n_xlsx += 1

        # 合并另存 HSF PDF
        path, w = save_pdf_group(self.out_dir, prefix, 'HSF', disp,
                                 block.hsf_objs, self._used_pdf_paths)
        warns.extend(w)
        if path:
            self.res.n_hsf += 1

        # MSDS：跟随 K 块，与输出的 HSF PDF 一一对应（2026-08-31 v1.3）
        group = next((g for (rs, re_, g) in msds_cells
                      if rs <= block.row_start <= re_), None)
        if group:
            mpath, mw = save_pdf_group(self.out_dir, prefix, 'MSDS', disp,
                                       group, self._used_pdf_paths)
            self.res.n_msds += 1
            rec.msds_out = mpath
            for x in mw:
                self.log(x, 'WARN')
            src = '、'.join(o.file_name for o in group)
            warns.append(f'MSDS 取自所属 M 列合并格（{src}），按本材料名另存')
        else:
            warns.append('所属 M 列合并格内没有 MSDS，未生成 MSDS PDF')

        self.log(f'{disp}（行 {block.excel_rows}）：成分 {len(rec.components)} 条，'
                 f'HSF {len(block.hsf_objs)} 份，测试项 {len(merged)} 项'
                 + (f'，告警 {len(warns)} 条' if warns else ''))
        return rec

    # ---------------- 日志 / 汇总
    def _write_log(self):
        path = os.path.join(self.out_dir, 'processing_log.txt')
        with open(path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(self._log_lines))
            f.write('\n')

    def _write_summary(self, prefix: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '汇总'
        head = ['序号', '材料名称', '源表行号', '输出文件', 'MSDS 文件', '成分数',
                'HSF 报告数', '测试机构', '报告编号', '报告日期', '告警']
        ws.append(head)
        for c in ws[1]:
            c.font = BOLD
        col_date = head.index('报告日期') + 1

        for rec in self.res.records:
            agencies = ' / '.join(sorted({r.agency for r in rec.reports if r.agency})) or ''
            nos = ' / '.join(r.report_no for r in rec.reports if r.report_no) or ''
            dates = ' / '.join(r.report_date for r in rec.reports if r.report_date) or ''
            row = [rec.index + 1, rec.name, rec.excel_rows,
                   os.path.basename(getattr(rec, 'out_xlsx', '') or ''),
                   os.path.basename(getattr(rec, 'msds_out', '') or ''),
                   len(rec.components), len(rec.reports),
                   agencies, nos, dates, '；'.join(rec.warnings)]
            ws.append(row)
            if rec.warnings:
                ws.cell(row=ws.max_row, column=len(head)).fill = RED_FILL
            # 整格日期都识别不出时留空（grill-me 决策），这里标红提醒人工补
            if rec.reports and not dates:
                ws.cell(row=ws.max_row, column=col_date).fill = RED_FILL

        # 明细：每份表的 29 项结果
        ws2 = wb.create_sheet('测试结果明细')
        ws2.append(['材料名称'] + [k for k, _ in RESULT_ITEMS])
        for c in ws2[1]:
            c.font = BOLD
        for rec in self.res.records:
            ws2.append([rec.name] + [rec.results.get(k, 'NA') for k, _ in RESULT_ITEMS])

        # 被跳过的空块：让用户一眼看到「哪些行没出表、为什么」
        ws3 = wb.create_sheet('跳过的空块')
        ws3.append(['源表行号', '跳过原因'])
        for c in ws3[1]:
            c.font = BOLD
        if self.res.skipped:
            for rs, re_, reason in self.res.skipped:
                ws3.append([f'{rs + 1}-{re_ + 1}', reason])
        else:
            ws3.append(['（无）', '本次运行没有空块，所有 K 列合并格都出了表'])

        for sheet in (ws, ws2, ws3):
            for col in sheet.columns:
                w = max((len(str(c.value)) for c in col if c.value is not None),
                        default=8)
                sheet.column_dimensions[col[0].column_letter].width = min(max(w + 2, 8), 60)

        path = os.path.join(self.out_dir, '_汇总报告.xlsx')
        wb.save(path)
